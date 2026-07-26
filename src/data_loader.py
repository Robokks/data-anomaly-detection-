"""Unified multi-format data loading (.tdms, .csv, .xlsx/.xls) and folder scanning.

Normalizes every supported format into the same shape used throughout the
pipeline: numeric channel columns, a ``time`` or ``sample`` index, and
``df.attrs["source_file"]`` set. This lets the rest of the codebase (feature
extraction, models, GUI) stay format-agnostic.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

from src.tdms_loader import load_tdms_dataframe

SUPPORTED_EXTENSIONS = (".tdms", ".csv", ".xlsx", ".xls")

_TIME_COLUMN_CANDIDATES = {
    "time", "timestamp", "t", "elapsed_time", "elapsed time", "time_s",
    "time (s)", "datetime", "date_time",
}


def infer_time_column(columns: list[str]) -> str | None:
    """Case-insensitive match of a column name against common time-column names."""
    lookup = {c.strip().lower(): c for c in columns}
    for candidate in _TIME_COLUMN_CANDIDATES:
        if candidate in lookup:
            return lookup[candidate]
    return None


def infer_sample_rate(df: pd.DataFrame) -> float | None:
    """Estimate sample rate in Hz from a numeric/datetime time index; None for a plain sample index."""
    if df.index.name != "time" or len(df) < 2:
        return None
    values = df.index.values
    if np.issubdtype(values.dtype, np.datetime64):
        diffs = np.diff(values).astype("timedelta64[ns]").astype(np.float64) / 1e9
    else:
        diffs = np.diff(values.astype(np.float64))
    median_dt = np.median(diffs)
    if median_dt <= 0:
        return None
    return float(1.0 / median_dt)


def _normalize_tabular(
    raw: pd.DataFrame,
    path: Path,
    fmt: str,
    time_column: str | None,
    channels: list[str] | None,
) -> pd.DataFrame:
    time_column = time_column or infer_time_column(list(raw.columns))

    df = raw.copy()
    index = None
    if time_column and time_column in df.columns:
        index = df[time_column]
        df = df.drop(columns=[time_column])

    for col in df.columns:
        df[col] = pd.to_numeric(df[col], errors="coerce")
    df = df.dropna(axis=1, how="all")

    if channels:
        keep = [c for c in df.columns if c in channels]
        df = df[keep]

    if df.shape[1] == 0:
        raise ValueError(f"No numeric channel columns found in {path}")

    if index is not None:
        try:
            index = pd.to_datetime(index)
        except (ValueError, TypeError):
            index = pd.to_numeric(index, errors="coerce")
        df.index = pd.Index(index.values, name="time")
        df = df[~pd.isna(df.index)]
    else:
        df.index = pd.RangeIndex(len(df), name="sample")

    df.attrs["source_file"] = str(path)
    df.attrs["format"] = fmt
    df.attrs["sample_rate_hz"] = infer_sample_rate(df)
    return df


def load_csv(
    path: str | Path,
    time_column: str | None = None,
    channels: list[str] | None = None,
) -> pd.DataFrame:
    path = Path(path)
    raw = pd.read_csv(path)
    return _normalize_tabular(raw, path, "csv", time_column, channels)


def load_excel(
    path: str | Path,
    sheet_name: str | int = 0,
    time_column: str | None = None,
    channels: list[str] | None = None,
) -> pd.DataFrame:
    path = Path(path)
    raw = pd.read_excel(path, sheet_name=sheet_name)
    return _normalize_tabular(raw, path, "excel", time_column, channels)


def load_dataframe(
    path: str | Path,
    group: str | None = None,
    channels: list[str] | None = None,
    time_column: str | None = None,
) -> pd.DataFrame:
    """Load a single .tdms/.csv/.xlsx/.xls file into a normalized DataFrame."""
    path = Path(path)
    ext = path.suffix.lower()

    if ext == ".tdms":
        df = load_tdms_dataframe(path, group_name=group, channels=channels)
        df.attrs["format"] = "tdms"
        df.attrs.setdefault("sample_rate_hz", infer_sample_rate(df))
        return df
    if ext == ".csv":
        return load_csv(path, time_column=time_column, channels=channels)
    if ext in (".xlsx", ".xls"):
        return load_excel(path, time_column=time_column, channels=channels)

    raise ValueError(f"Unsupported file extension '{ext}' for {path}")


def scan_directory(
    root: str | Path,
    recursive: bool = True,
    extensions: tuple[str, ...] = SUPPORTED_EXTENSIONS,
) -> list[Path]:
    """Find every supported data file under ``root``, regardless of naming or subdirectory depth."""
    root = Path(root)
    if not root.is_dir():
        raise NotADirectoryError(f"{root} is not a directory")

    walker = root.rglob("*") if recursive else root.glob("*")
    extensions_lower = {e.lower() for e in extensions}
    files = sorted(p for p in walker if p.is_file() and p.suffix.lower() in extensions_lower)
    return files


@dataclass
class FileInfo:
    path: Path
    extension: str
    size_bytes: int
    channels: list[str] = field(default_factory=list)
    n_samples: int | None = None
    duration_seconds: float | None = None
    sample_rate_hz: float | None = None
    error: str | None = None


def probe_file(path: str | Path) -> FileInfo:
    """Cheap metadata-only read: channel names and row count without loading full data where possible."""
    path = Path(path)
    ext = path.suffix.lower()
    info = FileInfo(path=path, extension=ext, size_bytes=path.stat().st_size)

    try:
        if ext == ".tdms":
            from nptdms import TdmsFile

            tdms_meta = TdmsFile.read_metadata(str(path))
            groups = tdms_meta.groups()
            if not groups:
                raise ValueError("no groups in TDMS file")
            group = groups[0]
            info.channels = [ch.name for ch in group.channels()]
            lengths = [len(ch) for ch in group.channels()]
            info.n_samples = int(min(lengths)) if lengths else None
        elif ext == ".csv":
            with open(path, "r", errors="replace") as f:
                header = f.readline()
                n_lines = sum(1 for _ in f)
            info.channels = [c.strip() for c in header.strip().split(",")]
            info.n_samples = n_lines
        elif ext in (".xlsx", ".xls"):
            import openpyxl

            wb = openpyxl.load_workbook(str(path), read_only=True)
            ws = wb[wb.sheetnames[0]]
            first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            info.channels = [str(c) for c in first_row if c is not None]
            info.n_samples = (ws.max_row or 1) - 1
            wb.close()
        else:
            raise ValueError(f"unsupported extension '{ext}'")
    except Exception as exc:  # noqa: BLE001 - surfaced to the GUI file list, not raised
        info.error = str(exc)

    return info


def scan_directory_with_info(
    root: str | Path,
    recursive: bool = True,
    extensions: tuple[str, ...] = SUPPORTED_EXTENSIONS,
) -> list[FileInfo]:
    """scan_directory() + probe_file() for every match."""
    return [probe_file(p) for p in scan_directory(root, recursive=recursive, extensions=extensions)]


def load_directory(
    directory: str | Path,
    recursive: bool = True,
    group: str | None = None,
    channels: list[str] | None = None,
) -> list[pd.DataFrame]:
    """Load every supported data file in a directory into a list of DataFrames."""
    files = scan_directory(directory, recursive=recursive)
    if not files:
        raise FileNotFoundError(f"No supported data files found in {directory}")
    return [load_dataframe(f, group=group, channels=channels) for f in files]
